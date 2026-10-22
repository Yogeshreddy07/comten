"""Check actual workbook cells and the fresh cloud workflow's business totals."""
import json,math,argparse
from pathlib import Path
from openpyxl import load_workbook
root=Path(__file__).resolve().parents[1]
parser=argparse.ArgumentParser();parser.add_argument('--android',action='store_true');args=parser.parse_args()
folder=root/('tests/exports/android' if args.android else 'tests/exports/cloud')
expected=json.loads((folder/'expected-cells.json').read_text())
cells=0
for item in expected:
    workbook=load_workbook(folder/item['name'],data_only=False)
    for sheet in item['sheets']:
        actual=workbook[sheet['name']]
        rows=[sheet['headers']]+sheet['rows']
        assert actual.max_row==len(rows),(item['name'],sheet['name'],'row count')
        for row_no,row in enumerate(rows,1):
            for col_no,value in enumerate(row,1):
                found=actual.cell(row_no,col_no)
                assert found.data_type!='f','Synthetic export cells must remain literal values'
                if isinstance(value,(int,float)) and not isinstance(value,bool): assert isinstance(found.value,(int,float)) and math.isclose(found.value,value,abs_tol=1e-9),(item['name'],row_no,col_no)
                else: assert (found.value or '')==(value or ''),(item['name'],row_no,col_no)
                cells+=1
journey=json.loads((root/'tests/android/device-results.json').read_text())['workflow'] if args.android else json.loads((root/'tests/cloud-journey.json').read_text())
code=journey['clientId']
expected_hours=round(journey.get('approvedSeconds',60)/3600,2)
def rows(prefix,sheet):
    item=next(i for i in expected if i['name'].startswith(prefix))
    return list(load_workbook(folder/item['name'],data_only=True)[sheet].values)[1:]
working=next(r for r in rows('working-hours','Working Hours') if r[0]==code)
assert working[2]==expected_hours and working[6]==expected_hours and working[7]==0
attendance=next(r for r in rows('attendance','Attendance') if r[0]==code)
assert attendance[2]==1
billing=rows('project-billing','Project Billing')[0]
assert billing[0]==code and billing[3]==expected_hours and billing[5]==expected_hours
utilization=next(r for r in rows('project-time-utilization','Projects') if r[1]==code)
assert utilization[3]==2.5 and utilization[4]==expected_hours and utilization[5]==0
client=next(r for r in rows('clients','Client-wise Projects') if r[0]==code)
assert client[1:]==(1,1,2.5,expected_hours)
status=next(r for r in rows('billing-status','Billing Status') if r[0]==code)
assert status[3]=='Billed' and status[4]==1 and status[5]==('Synthetic Android' if args.android else 'Synthetic')
result={'workbooks':len(expected),'cellsCompared':cells,'freshWorkflowFamilies':6,'passed':True}
(folder/'verification.json').write_text(json.dumps(result,indent=2))
print(json.dumps(result))
